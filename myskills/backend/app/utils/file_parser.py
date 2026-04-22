"""文件解析工具"""

import os
import docx
import pdfplumber
import openpyxl
import pandas as pd


def parse_md(file_path: str) -> str:
    """解析Markdown文件"""
    with open(file_path, 'r', encoding='utf-8') as f:
        return f.read()


def parse_txt(file_path: str) -> str:
    """解析文本文件"""
    with open(file_path, 'r', encoding='utf-8') as f:
        return f.read()


def parse_docx(file_path: str) -> str:
    """解析DOCX文件"""
    doc = docx.Document(file_path)
    text = []
    for paragraph in doc.paragraphs:
        text.append(paragraph.text)
    return '\n'.join(text)


def parse_pdf(file_path: str) -> str:
    """解析PDF文件"""
    text = []
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text.append(page_text)
    return '\n'.join(text)


def parse_excel(file_path: str) -> dict:
    """解析Excel文件"""
    workbook = openpyxl.load_workbook(file_path)
    sheets = {}
    
    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        data = []
        for row in sheet.iter_rows(values_only=True):
            # 过滤空行
            if any(cell is not None for cell in row):
                data.append(list(row))
        sheets[sheet_name] = data
    
    return sheets


def get_file_extension(file_path: str) -> str:
    """获取文件扩展名"""
    return os.path.splitext(file_path)[1].lower()[1:]


def parse_file(file_path: str) -> str:
    """根据文件类型解析文件"""
    ext = get_file_extension(file_path)
    
    if ext == 'md':
        return parse_md(file_path)
    elif ext == 'txt':
        return parse_txt(file_path)
    elif ext == 'docx':
        return parse_docx(file_path)
    elif ext == 'pdf':
        return parse_pdf(file_path)
    else:
        raise ValueError(f"不支持的文件类型: {ext}")
